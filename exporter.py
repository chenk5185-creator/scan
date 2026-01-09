"""
Excel exporter for wallet scan results.
Generates a 14-sheet Excel report.
"""
import logging
from datetime import datetime
from typing import Dict, List, Any, Optional
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

from processor import ProcessedData, DataProcessor
from config import SUPPORTED_CHAINS, CHAIN_TO_SHEET, DEFAULT_OUTPUT_FILE

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class ExcelExporter:
    """Export processed data to Excel with multiple sheets."""

    # Style definitions
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

    MONEY_FORMAT = '#,##0.00'
    PERCENT_FORMAT = '0.00%'
    NUMBER_FORMAT = '#,##0.00######'

    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    def __init__(self, processor: DataProcessor):
        """
        Initialize the exporter.

        Args:
            processor: DataProcessor with processed data
        """
        self.processor = processor
        self.workbook: Optional[Workbook] = None

    def export(self, output_path: str = DEFAULT_OUTPUT_FILE) -> str:
        """
        Export all data to Excel file.

        Args:
            output_path: Path for the output Excel file

        Returns:
            Path to the created file
        """
        if not self.processor.processed_data:
            raise ValueError("No processed data available. Run processor.process() first.")

        logger.info(f"Exporting data to {output_path}")
        self.workbook = Workbook()

        # Remove default sheet
        if "Sheet" in self.workbook.sheetnames:
            del self.workbook["Sheet"]

        # Create all 14 sheets
        self._create_summary_sheet()
        self._create_all_tokens_sheet()
        self._create_statistics_sheet()

        # Create chain-specific sheets
        for chain_id, sheet_name in CHAIN_TO_SHEET.items():
            self._create_chain_sheet(chain_id, sheet_name)

        self._create_errors_sheet()

        # Save workbook
        self.workbook.save(output_path)
        logger.info(f"Excel report saved to {output_path}")

        return output_path

    def _create_summary_sheet(self):
        """Create the Summary sheet with overall statistics."""
        ws = self.workbook.create_sheet("Summary")
        data = self.processor.processed_data

        # Title
        ws["A1"] = "Multi-Chain Wallet Scan Report"
        ws["A1"].font = Font(bold=True, size=16)
        ws.merge_cells("A1:E1")

        # Timestamp
        ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A2"].font = Font(italic=True)

        # Overall Stats
        row = 4
        ws[f"A{row}"] = "Overall Statistics"
        ws[f"A{row}"].font = Font(bold=True, size=14)
        row += 1

        stats = [
            ("Total Wallets Scanned", data.total_wallets),
            ("Successful Scans", data.successful_scans),
            ("Failed Scans", data.failed_scans),
            ("Success Rate", f"{(data.successful_scans / max(data.total_wallets, 1)) * 100:.1f}%"),
            ("Total Portfolio Value (USD)", f"${data.total_value_usd:,.2f}"),
            ("Total Token Holdings", data.total_tokens),
            ("Unique Tokens", len(data.token_summaries)),
        ]

        for label, value in stats:
            ws[f"A{row}"] = label
            ws[f"B{row}"] = value
            row += 1

        # Chain Distribution
        row += 1
        ws[f"A{row}"] = "Value Distribution by Chain"
        ws[f"A{row}"].font = Font(bold=True, size=14)
        row += 1

        # Headers
        headers = ["Chain", "Total Value (USD)", "Wallets", "Token Holdings", "% of Total"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
        row += 1

        # Chain data
        for chain_id, summary in sorted(
            data.chain_summaries.items(),
            key=lambda x: x[1].total_value_usd,
            reverse=True,
        ):
            if summary.total_value_usd > 0:
                pct = (summary.total_value_usd / max(data.total_value_usd, 1)) * 100
                ws.cell(row=row, column=1, value=summary.chain_name)
                ws.cell(row=row, column=2, value=summary.total_value_usd).number_format = self.MONEY_FORMAT
                ws.cell(row=row, column=3, value=summary.wallet_count)
                ws.cell(row=row, column=4, value=summary.token_count)
                ws.cell(row=row, column=5, value=f"{pct:.2f}%")
                row += 1

        # Top Wallets
        row += 2
        ws[f"A{row}"] = "Top 20 Wallets by Value"
        ws[f"A{row}"].font = Font(bold=True, size=14)
        row += 1

        headers = ["Rank", "Address", "Total Value (USD)", "Active Chains", "Token Count"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
        row += 1

        top_wallets = sorted(
            data.wallet_summaries.values(),
            key=lambda x: x.total_value_usd,
            reverse=True,
        )[:20]

        for rank, wallet in enumerate(top_wallets, 1):
            ws.cell(row=row, column=1, value=rank)
            ws.cell(row=row, column=2, value=wallet.address)
            ws.cell(row=row, column=3, value=wallet.total_value_usd).number_format = self.MONEY_FORMAT
            ws.cell(row=row, column=4, value=len(wallet.active_chains))
            ws.cell(row=row, column=5, value=wallet.token_count)
            row += 1

        # Auto-adjust column widths
        self._auto_adjust_columns(ws)

    def _create_all_tokens_sheet(self):
        """Create the All_Tokens sheet with all token holdings."""
        ws = self.workbook.create_sheet("All_Tokens")

        tokens = self.processor.get_all_tokens()
        if not tokens:
            ws["A1"] = "No token data available"
            return

        # Sort by value descending
        tokens = sorted(tokens, key=lambda x: x["value_usd"], reverse=True)

        # Create DataFrame
        df = pd.DataFrame(tokens)

        # Reorder columns
        columns = [
            "wallet_address",
            "chain_name",
            "token_symbol",
            "token_name",
            "token_address",
            "price_usd",
            "amount",
            "value_usd",
            "is_verified",
            "is_core",
        ]
        df = df[[c for c in columns if c in df.columns]]

        # Rename columns for display
        column_names = {
            "wallet_address": "Wallet Address",
            "chain_name": "Chain",
            "token_symbol": "Symbol",
            "token_name": "Token Name",
            "token_address": "Token Address",
            "price_usd": "Price (USD)",
            "amount": "Amount",
            "value_usd": "Value (USD)",
            "is_verified": "Verified",
            "is_core": "Core Token",
        }
        df = df.rename(columns=column_names)

        # Write to sheet
        self._write_dataframe(ws, df)

    def _create_statistics_sheet(self):
        """Create the Statistics sheet with detailed analytics."""
        ws = self.workbook.create_sheet("Statistics")
        data = self.processor.processed_data

        # Token Statistics
        row = 1
        ws[f"A{row}"] = "Token Statistics (Top 50 by Total Value)"
        ws[f"A{row}"].font = Font(bold=True, size=14)
        row += 1

        headers = ["Chain", "Symbol", "Token Name", "Address", "Total Value (USD)", "Total Amount", "Holders", "Price (USD)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
        row += 1

        top_tokens = sorted(
            data.token_summaries.values(),
            key=lambda x: x.total_value_usd,
            reverse=True,
        )[:50]

        for token in top_tokens:
            ws.cell(row=row, column=1, value=SUPPORTED_CHAINS.get(token.chain, token.chain))
            ws.cell(row=row, column=2, value=token.symbol)
            ws.cell(row=row, column=3, value=token.name)
            ws.cell(row=row, column=4, value=token.address)
            ws.cell(row=row, column=5, value=token.total_value_usd).number_format = self.MONEY_FORMAT
            ws.cell(row=row, column=6, value=token.total_amount).number_format = self.NUMBER_FORMAT
            ws.cell(row=row, column=7, value=token.holder_count)
            ws.cell(row=row, column=8, value=token.price).number_format = self.MONEY_FORMAT
            row += 1

        # Wallet Statistics
        row += 2
        ws[f"A{row}"] = "Wallet Value Distribution"
        ws[f"A{row}"].font = Font(bold=True, size=14)
        row += 1

        # Calculate distribution buckets
        values = [w.total_value_usd for w in data.wallet_summaries.values()]
        buckets = [
            ("$0", 0, 0.01),
            ("$0.01 - $10", 0.01, 10),
            ("$10 - $100", 10, 100),
            ("$100 - $1,000", 100, 1000),
            ("$1,000 - $10,000", 1000, 10000),
            ("$10,000 - $100,000", 10000, 100000),
            ("$100,000+", 100000, float('inf')),
        ]

        headers = ["Value Range", "Wallet Count", "Percentage"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
        row += 1

        total_wallets = len(values)
        for label, min_val, max_val in buckets:
            count = sum(1 for v in values if min_val <= v < max_val)
            pct = (count / max(total_wallets, 1)) * 100
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=count)
            ws.cell(row=row, column=3, value=f"{pct:.1f}%")
            row += 1

        # Chain Activity Statistics
        row += 2
        ws[f"A{row}"] = "Chain Activity Statistics"
        ws[f"A{row}"].font = Font(bold=True, size=14)
        row += 1

        headers = ["Active Chains", "Wallet Count", "Avg Value (USD)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
        row += 1

        # Group by number of active chains
        chain_count_groups: Dict[int, List[float]] = {}
        for wallet in data.wallet_summaries.values():
            n = len(wallet.active_chains)
            if n not in chain_count_groups:
                chain_count_groups[n] = []
            chain_count_groups[n].append(wallet.total_value_usd)

        for n in sorted(chain_count_groups.keys()):
            values = chain_count_groups[n]
            avg_value = sum(values) / len(values) if values else 0
            ws.cell(row=row, column=1, value=n)
            ws.cell(row=row, column=2, value=len(values))
            ws.cell(row=row, column=3, value=avg_value).number_format = self.MONEY_FORMAT
            row += 1

        self._auto_adjust_columns(ws)

    def _create_chain_sheet(self, chain_id: str, sheet_name: str):
        """Create a sheet for a specific chain."""
        ws = self.workbook.create_sheet(sheet_name)

        tokens = self.processor.get_chain_data(chain_id)
        if not tokens:
            ws["A1"] = f"No token data for {SUPPORTED_CHAINS.get(chain_id, chain_id)}"
            return

        # Sort by value descending
        tokens = sorted(tokens, key=lambda x: x["value_usd"], reverse=True)

        # Create DataFrame
        df = pd.DataFrame(tokens)

        # Select and rename columns
        columns = [
            "wallet_address",
            "token_symbol",
            "token_name",
            "token_address",
            "price_usd",
            "amount",
            "value_usd",
            "is_verified",
            "is_core",
        ]
        df = df[[c for c in columns if c in df.columns]]

        column_names = {
            "wallet_address": "Wallet Address",
            "token_symbol": "Symbol",
            "token_name": "Token Name",
            "token_address": "Token Address",
            "price_usd": "Price (USD)",
            "amount": "Amount",
            "value_usd": "Value (USD)",
            "is_verified": "Verified",
            "is_core": "Core Token",
        }
        df = df.rename(columns=column_names)

        # Write to sheet
        self._write_dataframe(ws, df)

    def _create_errors_sheet(self):
        """Create the Errors sheet with failed addresses."""
        ws = self.workbook.create_sheet("Errors")

        failed = self.processor.get_failed_addresses()
        if not failed:
            ws["A1"] = "No errors - all addresses scanned successfully!"
            ws["A1"].font = Font(color="008000")
            return

        # Create DataFrame
        df = pd.DataFrame(failed)
        df = df.rename(columns={"address": "Wallet Address", "error": "Error Message"})

        # Write to sheet
        self._write_dataframe(ws, df)

    def _write_dataframe(self, ws, df: pd.DataFrame):
        """Write a DataFrame to a worksheet with formatting."""
        # Write headers
        for col, header in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
            cell.border = self.THIN_BORDER

        # Write data
        for row_idx, row_data in enumerate(df.values, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)

                # Apply number format for numeric columns
                col_name = df.columns[col_idx - 1]
                if "Value" in col_name or "Price" in col_name:
                    cell.number_format = self.MONEY_FORMAT
                elif col_name == "Amount":
                    cell.number_format = self.NUMBER_FORMAT

        # Auto-adjust columns
        self._auto_adjust_columns(ws)

        # Freeze header row
        ws.freeze_panes = "A2"

    def _auto_adjust_columns(self, ws):
        """Auto-adjust column widths based on content."""
        for column_cells in ws.columns:
            max_length = 0
            column = column_cells[0].column_letter

            for cell in column_cells:
                try:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                except:
                    pass

            # Set width with some padding, but cap at reasonable max
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width


def export_to_excel(
    processor: DataProcessor,
    output_path: str = DEFAULT_OUTPUT_FILE,
) -> str:
    """
    Convenience function to export processed data to Excel.

    Args:
        processor: DataProcessor with processed data
        output_path: Output file path

    Returns:
        Path to created file
    """
    exporter = ExcelExporter(processor)
    return exporter.export(output_path)
